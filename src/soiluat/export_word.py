from __future__ import annotations

import re
import sqlite3
from pathlib import Path

from docx import Document

from soiluat import semantic
from soiluat.cards import OCR_LABEL, Card, _ocr_unverified
from soiluat.reader import cite

HIEU_LUC = {"con": "Còn hiệu lực", "het": "Hết hiệu lực", "chua_ro": "Hiệu lực chưa rõ"}  # như _results.html
_XML_CAM = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")  # Word (XML 1.0) không nhận; .txt kiểu DOS có ngắt trang \f, NUL


def _sach(text: str) -> str:
    return _XML_CAM.sub(" ", text)


def page_label(page_from, page_to=None) -> str:
    """"Trang x" hoặc "Trang x–y"; rỗng khi chưa biết trang."""
    if not page_from:
        return ""
    if page_to and page_to != page_from:
        return f"Trang {page_from}–{page_to}"
    return f"Trang {page_from}"


def card_pages(c: Card) -> str:
    # page_to / notes / intent do agent C thêm vào Card — đọc phòng thủ
    return page_label(getattr(c, "page_from", None), getattr(c, "page_to", None))


def export_units(conn: sqlite3.Connection, uids: list[int], dest: Path) -> None:
    """Báo cáo Word đúng các đơn vị đã chọn, nguyên văn (Điều kèm mọi khoản, điểm); không sinh câu nào."""
    doc = Document()
    doc.add_heading("Báo cáo Soi Luật", level=0)
    for uid in dict.fromkeys(uids):
        u = conn.execute(
            "SELECT u.doc_id, u.path, u.page_from, u.page_to, d.so_hieu, d.loai, d.file_path, d.hieu_luc "
            "FROM units u JOIN documents d ON d.id = u.doc_id WHERE u.id = ?", (uid,)
        ).fetchone()
        if u is None:  # đơn vị đã bị xoá sau khi kết quả hiện ra
            continue
        doc.add_heading(_sach(cite(u["so_hieu"], u["loai"], u["path"])), level=1)
        meta = [page_label(u["page_from"], u["page_to"]), u["file_path"], HIEU_LUC.get(u["hieu_luc"] or "chua_ro")]
        doc.add_paragraph(_sach(" · ".join(x for x in meta if x)))
        head, lines = semantic.root_lines(conn, uid)
        for t in [head, *(t for _, t in lines)]:
            if t:
                doc.add_paragraph(_sach(t))
        for (n,) in conn.execute(
            "SELECT COALESCE(reviewed_text, raw_text) FROM units WHERE parent_id = ? AND level = 'ghi_chu' "
            "ORDER BY id LIMIT 3", (uid,)
        ):
            doc.add_paragraph(_sach(f"Ghi chú: {n}"))
        if _ocr_unverified(conn, u["page_from"], u["page_to"], u["doc_id"]):
            doc.add_paragraph(f"Cảnh báo: {OCR_LABEL}")
    dest.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(dest))
