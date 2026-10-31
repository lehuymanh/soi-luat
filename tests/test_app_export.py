from pathlib import Path

from fastapi.testclient import TestClient

from soiluat.db import backup_db


def test_app_does_not_import_ocr():
    import sys

    sys.modules.pop("soiluat.ocr", None)
    sys.modules.pop("soiluat.app", None)
    import soiluat.app  # noqa: F401

    assert "soiluat.ocr" not in sys.modules


def test_http_search(tmp_path, monkeypatch):
    monkeypatch.setenv("SOILUAT_DATA", str(tmp_path))
    from soiluat.app import app
    from soiluat.config import db_path
    from soiluat.db import connect, init_schema
    from soiluat.ingest import ingest_path

    conn = connect(db_path())
    init_schema(conn)
    src = Path(__file__).parent / "fixtures" / "nd15.txt"
    dest = tmp_path / "nd15.txt"
    dest.write_text(src.read_text(encoding="utf-8"), encoding="utf-8")
    ingest_path(conn, dest, tmp_path)
    conn.close()

    client = TestClient(app)
    r = client.get("/", params={"q": "bao cao dinh ky"})
    assert r.status_code == 200
    assert "Soi Luật" in r.text
    assert "Điều" in r.text or "báo cáo" in r.text.lower()
    assert "Hiệu lực chưa rõ" in r.text


def _doc77(db, tmp_path):
    """Nghị định có Điều 5 gồm ba khoản (khoản 3 có hai điểm) và Điều 12 phạt tiền."""
    text = (Path(__file__).parent / "fixtures" / "nd15.txt").read_text(encoding="utf-8")
    text = text.replace("15/2020/NĐ-CP", "77/2021/NĐ-CP").replace(
        "trước ngày 15 tháng đầu tiên của quý.",
        "trước ngày 15 tháng đầu tiên của quý.\n2. Báo cáo được gửi bằng dữ liệu điện tử.\n"
        "3. Hồ sơ báo cáo được lưu trữ năm năm.\na) bản chính;\nb) bản sao.",
    )
    from soiluat.ingest import ingest_path

    dest = tmp_path / "nd77.txt"
    dest.write_text(text, encoding="utf-8")
    return ingest_path(db, dest, tmp_path)


def _uid(db, doc_id, path):
    return db.execute("SELECT id FROM units WHERE doc_id = ? AND path = ?", (doc_id, path)).fetchone()["id"]


def _export_text(monkeypatch, tmp_path, uids) -> str:
    from docx import Document

    monkeypatch.setenv("SOILUAT_DATA", str(tmp_path))
    from soiluat.app import app

    r = TestClient(app).post("/export", data={"uid": uids})
    assert r.status_code == 200
    out = tmp_path / "got.docx"
    out.write_bytes(r.content)
    return "\n".join(p.text for p in Document(str(out)).paragraphs)


def test_c2_export_la_nguyen_van_dieu_da_chon_khong_co_cau_may_goi_y(db, tmp_path, monkeypatch):
    """C2: ô được tick xuất đúng đơn vị đó (tiêu đề trích dẫn + nguyên văn), kể cả khi nó có liên kết 'goi_y'."""
    doc = _doc77(db, tmp_path)
    phat, k2 = _uid(db, doc, "D12.K1"), _uid(db, doc, "D5.K2")
    db.execute("INSERT INTO links(sanction_unit, duty_unit, method, confirmed) VALUES (?, ?, 'goi_y', 0)", (phat, k2))
    db.commit()
    text = _export_text(monkeypatch, tmp_path, [k2])
    assert "Khoản 2 Điều 5 Nghị định 77/2021/NĐ-CP" in text
    assert "Báo cáo được gửi bằng dữ liệu điện tử." in text
    assert "nd77.txt" in text and "Hiệu lực chưa rõ" in text
    assert "có thể bị xử phạt theo" not in text and "Điều 12" not in text and "Hạng" not in text


def test_c2_tick_ca_dieu_xuat_du_cac_khoan_va_diem(db, tmp_path, monkeypatch):
    """C2: tick gốc Điều 5 thì Word có cả khoản 1, 2, 3 và điểm a, b, không chỉ khoản đầu."""
    doc = _doc77(db, tmp_path)
    text = _export_text(monkeypatch, tmp_path, [_uid(db, doc, "D5")])
    assert "Điều 5 Nghị định 77/2021/NĐ-CP" in text and "Báo cáo định kỳ" in text
    for must in ("trước ngày 15 tháng đầu tiên", "dữ liệu điện tử", "lưu trữ năm năm", "bản chính", "bản sao"):
        assert must in text, must


def test_export_and_backup(db, sample_doc, tmp_path):
    from docx import Document

    from soiluat.export_word import export_units

    dest = tmp_path / "out.docx"
    ids = [r["id"] for r in db.execute("SELECT id FROM units WHERE level = 'dieu' ORDER BY id LIMIT 2")]
    export_units(db, ids + [10**9], dest)  # id không còn trong kho thì bỏ qua
    assert "Điều 1" in "\n".join(p.text for p in Document(str(dest)).paragraphs)
    bak = tmp_path / "backup" / "soiluat.db"
    backup_db(tmp_path / "soiluat.db", bak)
    assert bak.exists() and bak.stat().st_size > 0
    from soiluat.db import connect

    restored = connect(bak)
    n = restored.execute("SELECT COUNT(*) AS n FROM units").fetchone()["n"]
    restored.close()
    assert n > 0
