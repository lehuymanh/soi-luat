from pathlib import Path

from docx import Document

from soiluat.extract import extract_docx
from soiluat.ingest import ingest_path


def test_ingest_docx(db, tmp_path):
    d = Document()
    d.add_paragraph("Nghị định số 99/2021/NĐ-CP")
    d.add_paragraph("Điều 2. Nghĩa vụ báo cáo")
    d.add_paragraph("1. Doanh nghiệp phải gửi báo cáo định kỳ.")
    path = tmp_path / "a.docx"
    d.save(str(path))
    doc_id = ingest_path(db, path, tmp_path)
    n = db.execute(
        "SELECT COUNT(*) AS n FROM units WHERE doc_id = ?", (doc_id,)
    ).fetchone()["n"]
    assert n >= 2
    labs = db.execute(
        "SELECT COUNT(*) AS n FROM unit_labels WHERE label = 'NGHIA_VU'"
    ).fetchone()["n"]
    assert labs >= 1


def test_docx_list_number_prefix(tmp_path):
    d = Document()
    d.add_paragraph("Điều 4. Báo cáo")
    d.add_paragraph("Doanh nghiệp phải gửi báo cáo định kỳ.", style="List Number")
    path = tmp_path / "n.docx"
    d.save(str(path))
    text = extract_docx(path)[0].text
    assert "Điều 4" in text
    assert "1." in text or "1)" in text


def _docx_xml_lon(path: Path, n_doan: int) -> None:
    """Docx hợp lệ mà document.xml có n_doan đoạn rỗng: nén còn rất nhỏ, giải nén ra rất lớn."""
    import zipfile

    Document().save(str(path))
    with zipfile.ZipFile(path) as z:
        parts = {n: z.read(n) for n in z.namelist()}
    parts["word/document.xml"] = parts["word/document.xml"].replace(b"<w:body>", b"<w:body>" + b"<w:p/>" * n_doan, 1)
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for n, b in parts.items():
            z.writestr(n, b)


def test_sec3_docx_giai_nen_qua_lon_bi_tu_choi_truoc_khi_parse(tmp_path, monkeypatch):
    """SEC-3: file docx vài chục KB mà XML giải nén hàng trăm MB làm luồng nạp (duy nhất) treo hàng chục phút: từ chối sớm."""
    import pytest

    from soiluat import extract

    monkeypatch.setattr(extract, "MAX_XML_BYTES", 100_000)
    path = tmp_path / "bom.docx"
    _docx_xml_lon(path, 40_000)  # ~240 KB XML
    assert path.stat().st_size < 50_000  # nén còn bé hơn nhiều so với 240 KB XML
    with pytest.raises(ValueError, match="quá lớn"):
        extract_docx(path)


def test_sec3_docx_qua_nhieu_doan_bi_tu_choi(tmp_path, monkeypatch):
    """SEC-3: dưới trần byte vẫn có thể có hàng triệu đoạn rỗng (~0,5 ms mỗi đoạn): cũng có trần số đoạn."""
    import pytest

    from soiluat import extract

    monkeypatch.setattr(extract, "MAX_PARAGRAPHS", 100)
    path = tmp_path / "doan.docx"
    _docx_xml_lon(path, 500)
    with pytest.raises(ValueError, match="quá nhiều đoạn"):
        extract_docx(path)


def test_sec3_xlsx_giai_nen_qua_lon_bi_tu_choi(tmp_path, monkeypatch):
    import zipfile

    import pytest

    from soiluat import extract

    monkeypatch.setattr(extract, "MAX_XML_BYTES", 100_000)
    path = tmp_path / "bom.xlsx"
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("xl/workbook.xml", "<workbook>" + " " * 200_000 + "</workbook>")
    with pytest.raises(ValueError, match="quá lớn"):
        extract.extract_xlsx(path)
